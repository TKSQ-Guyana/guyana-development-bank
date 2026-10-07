"""Citizens from the MPS call list — the workbook MPS shares with GDB every day.

MPS phones people who asked about a GDB loan and writes down who they are:
name, address, region, phone and an ID number, one sheet per calling round.
The platform administrator uploads that workbook here, and every person in it
who does not yet have a GDB account gets one — a citizen account in the
citizen realm, signed in through the National ID door (tin_auth.py), exactly
as if they had signed up online.

TWO STEPS, NOTHING CREATED BY THE FIRST.

  preview(actor, file_name, content)
      Reads every sheet, cleans every row and sorts it: New, Existing (an
      account already holds this ID), Duplicate (the same ID earlier in the
      file) or Error (with the reason). Writes one GDB Citizen Import holding
      the rows, the uploaded workbook and the error report — and no account.
  run(actor, batch, reason)
      Queues the batch to the long worker (execute), which opens an account
      for each New row. A batch runs once: asking again answers its state.

The file is cumulative — tomorrow's workbook still holds today's sheets — so
"what is new" is decided against the accounts GDB holds, never against the
file alone, and is decided again row by row at the moment of creation.

CLEANING RULES (GDB, 2026-10-07):
  * Name: the last word is the last name, every word before it the first name.
    One word, or anything but letters, apostrophes, dots and hyphens, is an
    error — a phone number typed into the name cell must not become a name.
  * ID: begins with a letter -> Passport; nine digits -> National ID; eleven
    digits -> e-ID; anything else ("NA", eight or ten digits, blank) is an
    error. Kept compact and upper-case: it is the Keycloak username.
  * Phone: the first number when two are given, as +592 and seven digits. A
    row without one is an error — the temporary password goes to it by SMS,
    and every later sign-in code too.
  * Region: the number MPS writes, mapped onto the Region list GDB holds. One
    that is not on it is left blank with a warning; it does not stop the row.
  * Business, type of business, amount and outcome are not imported.

EACH ACCOUNT, as complete_signup makes one: the Keycloak account (ID number as
username, placeholder email under tin_auth.PLACEHOLDER_DOMAIN, a TEMPORARY
password), the Frappe Website User holding Citizen, and the declared half of
GDB Citizen Profile — marked `declared_source`, because MPS wrote it and the
applicant did not. One row is one transaction: if the portal half fails the
Keycloak account is deleted again, and one bad row never stops the rest.

THE PASSWORD exists in exactly two places: Keycloak, as temporary, and the
text to the person's phone. It is not stored, logged, queued or returned. The
person replaces it at their first sign-in (tin_auth.set_initial_password); if
the text never arrived, "Forgot password" sends a code to the same phone and
sets a permanent one, which clears the temporary state too.
"""

import hashlib
import io
import json
import re
import secrets
import zipfile
from contextlib import contextmanager

import frappe
from frappe import _
from frappe.utils import cint, now_datetime

from gdb_bank import profiles, tin_auth
from gdb_bank.install import REGION_OPTIONS
from gdb_bank.integrations import keycloak_admin, sms
from gdb_bank.services import access_audit
from gdb_bank.utils.session import _logger

DOCTYPE = "GDB Citizen Import"
ROW_DOCTYPE = "GDB Citizen Import Row"
ID_TYPE_FIELD = "gdb_id_type"

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
# How far down a sheet the header row may sit (the sample has a title and a
# date above it).
HEADER_SCAN = 10
NAME_MAX = tin_auth.NAME_MAX
ADDRESS_MAX = 500
SMS_MAX = 150

PREVIEWED, QUEUED, RUNNING, COMPLETED, FAILED = "Previewed", "Queued", "Running", "Completed", "Failed"
NEW, EXISTING, DUPLICATE, ERROR, CREATED, ROW_FAILED = "New", "Existing", "Duplicate", "Error", "Created", "Failed"
NATIONAL_ID, PASSPORT, EID = "National ID", "Passport", "e-ID"

# Header text as MPS writes it (lower-cased, spaces collapsed) -> our key. Only
# these columns are read; matching by name, not position, because the sheets
# do not keep one column order.
HEADERS = {
	"name": "name",
	"full name": "name",
	"address": "address",
	"region": "region",
	"contact #": "phone",
	"contact": "phone",
	"contact no": "phone",
	"contact number": "phone",
	"phone": "phone",
	"id number": "id_number",
	"id no": "id_number",
	"id": "id_number",
}
REQUIRED = ("name", "address", "region", "phone", "id_number")
REQUIRED_LABELS = "NAME, ADDRESS, REGION, CONTACT # and ID Number"

NAME_PART = re.compile(r"^[A-Za-z][A-Za-z'.\-]*$")
PASSPORT_SHAPE = re.compile(r"^[A-Z]{1,3}\d{4,12}$")
REGIONS = [option for option in REGION_OPTIONS.split("\n") if option]

# The text that carries the temporary password (GDB, 2026-10-07). At most
# SMS_MAX characters once filled; welcome_text shortens the name to fit.
WELCOME_TEXT = (
	"Hi {first_name}, your GDB account is created via MPSGEI. Apply for a loan at "
	"gdb.gov.gy with ID {id_number} & temp password {password}. Thanks-GDB Team"
)

# Typed from a text message, so nothing that reads as something else (0/O, 1/l/I).
_PASSWORD_ALPHABET = "ABCDEFGHJKMNPQRSTUVWXYZabcdefghjkmnpqrstuvwxyz23456789"
PASSWORD_LENGTH = 8


class RowError(ValueError):
	"""A row that cannot become an account, and why — said to the administrator."""


# --------------------------------------------------------------------------
# cleaning — plain functions of one cell each
# --------------------------------------------------------------------------


def _text(value) -> str:
	"""A cell as text: Excel's 129093615.0 is "129093615", spaces collapsed."""
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	text = re.sub(r"\s+", " ", str(value)).strip()
	return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def split_name(value) -> tuple[str, str]:
	"""(first name, last name): the last word is the last name."""
	text = _text(value)
	if not text:
		raise RowError("Name is blank.")
	parts = text.split(" ")
	if not all(NAME_PART.match(part) for part in parts):
		raise RowError("Name has digits or symbols in it — only the person's name belongs there.")
	if len(parts) < 2:
		raise RowError("Name is one word — a first and a last name are needed.")
	first, last = " ".join(parts[:-1]), parts[-1]
	if len(first) > NAME_MAX or len(last) > NAME_MAX:
		raise RowError(f"Name is longer than {NAME_MAX} characters.")
	return first, last


def classify_id(value) -> tuple[str, str]:
	"""(ID type, ID number as GDB keeps it)."""
	compact = re.sub(r"[\s\-/.]", "", _text(value).upper())
	if not compact:
		raise RowError("ID number is blank.")
	if compact[0].isalpha():
		if not PASSPORT_SHAPE.match(compact) or len(compact) > 15:
			raise RowError(f"ID number {compact} begins with a letter but is not a passport number.")
		return PASSPORT, compact
	if re.fullmatch(r"\d{9}", compact):
		return NATIONAL_ID, compact
	if re.fullmatch(r"\d{11}", compact):
		return EID, compact
	raise RowError(
		f"ID number {compact} is not a 9-digit National ID, an 11-digit e-ID or a passport number."
	)


def clean_phone(value) -> str:
	"""The first number in the cell, as +592 and seven digits."""
	text = _text(value)
	if not text:
		raise RowError("Contact number is blank.")
	first = _text(re.split(r"/|,|;|\bor\b", text)[0])
	number = sms.guyana_number(first)
	if not number:
		raise RowError(f"Contact number {text} is not a Guyana phone number (seven digits).")
	return number


def clean_region(value) -> tuple[str | None, str | None]:
	"""(the Region option, or None; a warning, or None)."""
	text = _text(value)
	if not text:
		return None, "Region is blank."
	match = re.fullmatch(r"(?:region\s*)?(\d{1,2})", text, re.IGNORECASE)
	if match:
		prefix = f"Region {int(match.group(1))} "
		for option in REGIONS:
			if option.startswith(prefix):
				return option, None
	for option in REGIONS:
		if text.lower() == option.lower():
			return option, None
	return None, f"Region {text} is not one of GDB's regions — left blank."


def clean_address(value) -> tuple[str | None, str | None]:
	text = _text(value)[:ADDRESS_MAX]
	return (text, None) if text else (None, "Address is blank.")


def clean_row(values: dict) -> dict:
	"""One row's cells -> what an account would be made of, plus its verdict.
	Every problem in the row is reported, not just the first."""
	out = {"full_name": _text(values.get("name"))[:140], "result": NEW}
	errors, warnings = [], []
	try:
		out["first_name"], out["last_name"] = split_name(values.get("name"))
	except RowError as exc:
		errors.append(str(exc))
	try:
		out["id_type"], out["id_number"] = classify_id(values.get("id_number"))
	except RowError as exc:
		errors.append(str(exc))
		out["id_number"] = _text(values.get("id_number"))[:40] or None
	try:
		out["phone"] = clean_phone(values.get("phone"))
	except RowError as exc:
		errors.append(str(exc))
	out["region"], warning = clean_region(values.get("region"))
	if warning:
		warnings.append(warning)
	out["address"], warning = clean_address(values.get("address"))
	if warning:
		warnings.append(warning)
	if errors:
		out["result"] = ERROR
	out["message"] = " ".join(errors + warnings) or None
	return out


# --------------------------------------------------------------------------
# reading the workbook
# --------------------------------------------------------------------------


def _header_key(value) -> str | None:
	text = re.sub(r"\s+", " ", str(value or "")).strip().lower()
	return HEADERS.get(text)


def _plain(value):
	"""A cell as JSON can hold it, the way the sheet showed it."""
	if value is None or isinstance(value, bool | int | str):
		return value
	if isinstance(value, float):
		return int(value) if value.is_integer() else value
	return str(value)


def _is_xlsx(content: bytes) -> bool:
	"""By what the file is, not what it is called: an xlsx is a zip holding a
	workbook part."""
	if not content.startswith(b"PK\x03\x04"):
		return False
	try:
		with zipfile.ZipFile(io.BytesIO(content)) as archive:
			return "xl/workbook.xml" in archive.namelist()
	except zipfile.BadZipFile:
		return False


def read_workbook(content: bytes) -> tuple[list[dict], list[str]]:
	"""(rows, notes). Each row: sheet, row_number, values (our keys) and source
	(every column as the sheet headed it). Notes say which sheets were skipped."""
	from openpyxl import load_workbook

	if len(content) > MAX_BYTES:
		frappe.throw(_("The workbook is larger than {0} MB.").format(MAX_BYTES // 1024 // 1024))
	if not _is_xlsx(content):
		frappe.throw(_("This is not an Excel .xlsx workbook."))
	try:
		workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	except Exception as exc:
		_logger().info(f"citizen import: workbook unreadable ({type(exc).__name__})")
		frappe.throw(_("The workbook could not be read. Save it again from Excel and upload it."))

	rows, notes = [], []
	try:
		for sheet in workbook.worksheets:
			if sheet.sheet_state != "visible":
				notes.append(f"Sheet {sheet.title}: hidden — skipped.")
				continue
			columns = header = None
			for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
				if columns is None:
					if index > HEADER_SCAN:
						break
					keys = {_header_key(v): i for i, v in reversed(list(enumerate(values))) if _header_key(v)}
					if "name" not in keys:
						continue
					missing = [key for key in REQUIRED if key not in keys]
					if missing:
						notes.append(
							f"Sheet {sheet.title}: the header on row {index} has no "
							f"{', '.join(missing).replace('_', ' ')} column — skipped."
						)
						break
					columns, header = keys, values
					continue
				if not any(_text(v) for v in values):
					continue
				rows.append(
					{
						"sheet": sheet.title[:140],
						"row_number": index,
						"values": {key: values[i] if i < len(values) else None for key, i in columns.items()},
						"source": {
							_text(h): _plain(values[i])
							for i, h in enumerate(header)
							if _text(h) and i < len(values)
						},
					}
				)
				if len(rows) > MAX_ROWS:
					frappe.throw(_("The workbook has more than {0} rows. Split it and upload each part.").format(MAX_ROWS))
			if columns is None and not any(n.startswith(f"Sheet {sheet.title}:") for n in notes):
				notes.append(
					f"Sheet {sheet.title}: no header row with {REQUIRED_LABELS} in its first {HEADER_SCAN} rows — skipped."
				)
	finally:
		workbook.close()
	return rows, notes


# --------------------------------------------------------------------------
# what is new
# --------------------------------------------------------------------------


def _formatted_eid(number: str) -> str:
	return f"{number[:3]}-{number[3:7]}-{number[7:]}"


def _holders(ids: list[str]) -> dict[str, str]:
	"""ID number -> how GDB already knows it, for the IDs some account holds."""
	if not ids:
		return {}
	found = {}
	for field in (tin_auth.NID_FIELD, tin_auth.TIN_FIELD):
		for number in frappe.get_all("User", filters={field: ["in", ids]}, pluck=field):
			found[number] = "An account already holds this ID number."
	eids = {_formatted_eid(n): n for n in ids if re.fullmatch(r"\d{11}", n)}
	if eids:
		for field in ("gdb_eid", "gdb_staff_eid"):
			for value in frappe.get_all("User", filters={field: ["in", list(eids)]}, pluck=field):
				found[eids[value]] = "An account already signs in with this e-ID."
	for number in frappe.get_all(profiles.DOCTYPE, filters={"national_id": ["in", ids]}, pluck="national_id"):
		found.setdefault(number, "A citizen has already given this ID number on their profile.")
	return found


def sort_rows(cleaned: list[dict]) -> None:
	"""Mark duplicates within the file and IDs GDB already holds, in place."""
	seen = {}
	for row in cleaned:
		if row["result"] != NEW:
			continue
		first = seen.get(row["id_number"])
		if first:
			row["result"] = DUPLICATE
			row["message"] = f"Same ID number as sheet {first['sheet']}, row {first['row_number']}."
		else:
			seen[row["id_number"]] = row
	held = _holders(list(seen))
	for number, why in held.items():
		row = seen[number]
		row["result"] = EXISTING
		row["message"] = why


# --------------------------------------------------------------------------
# preview
# --------------------------------------------------------------------------


def _attach(doc, file_name: str, content: bytes) -> str:
	stored = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": file_name,
			"attached_to_doctype": DOCTYPE,
			"attached_to_name": doc.name,
			"is_private": 1,
			"content": content,
		}
	).insert(ignore_permissions=True)
	return stored.file_url


def _safe_file_name(name: str | None) -> str:
	base = re.sub(r"[^\w.\- ]", "_", (name or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]).strip() or "workbook"
	return base if base.lower().endswith(".xlsx") else f"{base}.xlsx"


def preview(actor: str, file_name: str | None, content: bytes) -> dict:
	rows, notes = read_workbook(content)
	if not rows:
		frappe.throw(
			_("No rows were found. Each sheet needs a header row with {0}.").format(REQUIRED_LABELS)
		)
	cleaned = []
	for raw in rows:
		row = clean_row(raw["values"])
		row.update(sheet=raw["sheet"], row_number=raw["row_number"], source=json.dumps(raw["source"], default=str))
		cleaned.append(row)
	sort_rows(cleaned)

	file_name = _safe_file_name(file_name)
	doc = frappe.get_doc(
		{
			"doctype": DOCTYPE,
			"status": PREVIEWED,
			"source_file_name": file_name,
			"file_hash": hashlib.sha256(content).hexdigest(),
			"uploaded_by": actor,
			"uploaded_on": now_datetime(),
			"notes": "\n".join(notes) or None,
			"rows": cleaned,
		}
	)
	_count(doc)
	doc.insert(ignore_permissions=True)
	doc.db_set("source_file", _attach(doc, file_name, content))
	_write_report(doc)
	frappe.db.commit()
	_logger().info(f"citizen import {doc.name} previewed by {actor}: {doc.total_rows} rows, {doc.new_rows} new")
	return summary(doc)


# --------------------------------------------------------------------------
# run
# --------------------------------------------------------------------------


def _load(batch: str):
	if not batch or not frappe.db.exists(DOCTYPE, batch):
		frappe.throw(_("No such import."), frappe.DoesNotExistError)
	return frappe.get_doc(DOCTYPE, batch)


def run(actor: str, batch: str, reason: str) -> dict:
	"""Queue the batch. Idempotent: a batch already queued, running or done is
	answered as it stands, never queued twice."""
	reason = (reason or "").strip()
	if len(reason) < 3:
		frappe.throw(_("Say why these accounts are being created."))
	doc = _load(batch)
	if doc.status != PREVIEWED:
		return summary(doc)
	if not doc.new_rows:
		frappe.throw(_("Nothing in this file is new — there are no accounts to create."))
	if not keycloak_admin.citizen_config():
		frappe.throw(_("Citizen account management in Keycloak is not configured on this site."))

	doc.status = QUEUED
	doc.run_by = actor
	doc.reason = reason[:500]
	doc.save(ignore_permissions=True)
	frappe.enqueue(
		"gdb_bank.services.citizen_import.execute",
		queue="long",
		timeout=3600,
		job_id=f"gdb_citizen_import::{doc.name}",
		deduplicate=True,
		enqueue_after_commit=True,
		batch=doc.name,
	)
	frappe.db.commit()
	_logger().info(f"citizen import {doc.name} queued by {actor}: {doc.new_rows} accounts")
	return summary(doc)


def execute(batch: str) -> None:
	"""The background job: an account for every New row, one at a time."""
	doc = frappe.get_doc(DOCTYPE, batch)
	if doc.status != QUEUED:
		return
	doc.status = RUNNING
	doc.started_on = now_datetime()
	doc.save(ignore_permissions=True)
	frappe.db.commit()

	try:
		with _room_for_new_users():
			_open_accounts(doc)
		doc.reload()
		_count(doc)
		doc.status = COMPLETED
	except Exception as exc:
		frappe.db.rollback()
		doc.reload()
		_count(doc)
		doc.status = FAILED
		doc.failure = f"{type(exc).__name__}: the import stopped. Rows already marked Created have their accounts."
		_logger().error(f"citizen import {batch} stopped: {type(exc).__name__}: {exc}")
	doc.finished_on = now_datetime()
	doc.save(ignore_permissions=True)
	_write_report(doc)
	frappe.db.commit()
	_logger().info(f"citizen import {batch} {doc.status.lower()}: {doc.created_count} created, {doc.failed_count} failed")


def _open_accounts(doc) -> None:
	for row in doc.rows:
		if row.result != NEW:
			continue
		result, message, user = _open_account(doc, row)
		# Bookkeeping on the import's own log, one row at a time so the
		# administrator's screen can follow along. The account itself was
		# saved through the ORM in _open_account.
		frappe.db.set_value(ROW_DOCTYPE, row.name, {"result": result, "message": message, "user": user})
		row.result, row.message, row.user = result, message, user
		_count(doc)
		frappe.db.set_value(
			DOCTYPE,
			doc.name,
			{"created_count": doc.created_count, "failed_count": doc.failed_count},
			update_modified=False,
		)
		frappe.db.commit()


def temporary_password() -> str:
	while True:
		chars = [secrets.choice(_PASSWORD_ALPHABET) for _ in range(PASSWORD_LENGTH)]
		if any(c.isupper() for c in chars) and any(c.islower() for c in chars) and any(c.isdigit() for c in chars):
			return "".join(chars)


def welcome_text(first_name: str, id_number: str, password: str) -> str:
	"""The text, at most SMS_MAX characters: the first word of the first name,
	shortened if the ID is long enough to need the room."""
	name = (first_name or "").split(" ")[0]
	room = SMS_MAX - len(WELCOME_TEXT.format(first_name="", id_number=id_number, password=password))
	return WELCOME_TEXT.format(first_name=name[: max(room, 0)], id_number=id_number, password=password)


# Frappe refuses a new User ("Throttled") once more than `throttle_user_limit`
# (default 60) were created in the last 60 MINUTES — its guard against sign-up
# abuse (user.throttle_user_creation). A day's call list is hundreds. The limit
# is raised by this many for the import job's own worker process only, and put
# back when the job ends: the web processes serving online sign-up keep the
# guard as configured.
THROTTLE_HEADROOM = MAX_ROWS


@contextmanager
def _room_for_new_users():
	conf = frappe.local.conf
	had = "throttle_user_limit" in conf
	before = conf.get("throttle_user_limit")
	conf["throttle_user_limit"] = cint(before or 60) + THROTTLE_HEADROOM
	try:
		yield
	finally:
		if had:
			conf["throttle_user_limit"] = before
		else:
			conf.pop("throttle_user_limit", None)


def _open_account(batch, row) -> tuple[str, str | None, str | None]:
	"""(result, message, user) for one New row."""
	number = row.id_number
	username = tin_auth._username(number)
	email = f"{username}@{tin_auth.PLACEHOLDER_DOMAIN}"
	# Decided again now: an account may have been opened since the preview.
	if tin_auth._account_for(number) or frappe.db.exists("User", email):
		return EXISTING, "An account was opened for this ID number after the preview.", None

	password = temporary_password()
	try:
		kc_id = keycloak_admin.create_citizen_account(
			username, email, row.first_name, row.last_name, password, temporary=True
		)
	except keycloak_admin.AccountExists:
		return EXISTING, "Keycloak already has a sign-in account for this ID number.", None
	except keycloak_admin.PasswordRejected:
		return ROW_FAILED, "Keycloak refused the temporary password (the realm's password policy).", None
	except keycloak_admin.KeycloakAdminError as exc:
		return ROW_FAILED, str(exc), None

	try:
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": email,
				"first_name": row.first_name,
				"last_name": row.last_name,
				"mobile_no": row.phone,
				"user_type": "Website User",
				"send_welcome_email": 0,
				"enabled": 1,
				tin_auth.NID_FIELD: number,
				ID_TYPE_FIELD: row.id_type,
				# Citizen and only Citizen — as the online sign-up provisions.
				"roles": [{"role": "Citizen"}],
			}
		).insert(ignore_permissions=True)
		profile = frappe.get_doc(profiles.DOCTYPE, profiles._ensure(user.name))
		profile.phone = row.phone
		profile.national_id = number
		profile.region = row.region or profile.region
		profile.address = row.address or profile.address
		profile.declared_source = f"MPS call list — {batch.name}, sheet {row.sheet}, row {row.row_number}"[:140]
		profile.save(ignore_permissions=True)
		reason = f"MPS call list {batch.name}: {batch.reason}"[:500]
		access_audit.record(
			batch.run_by, access_audit.ACCOUNT_CREATED, reason=reason, subject=user.name, subject_user=user.name, new=["Citizen"]
		)
		access_audit.record(
			batch.run_by, access_audit.ONE_TIME_PASSWORD_ISSUED, reason=reason, subject=user.name, subject_user=user.name
		)
		frappe.db.commit()
	except Exception as exc:
		# The Keycloak half must not outlive a portal account that was never made.
		frappe.db.rollback()
		try:
			keycloak_admin.delete_citizen_account(kc_id)
		except keycloak_admin.KeycloakAdminError:
			_logger().error(f"citizen import {batch.name}: Keycloak account {kc_id} left behind for row {row.idx}")
		_logger().error(f"citizen import {batch.name} row {row.idx}: {type(exc).__name__}: {exc}")
		message = str(exc) if isinstance(exc, frappe.ValidationError) else _("The portal account could not be created.")
		return ROW_FAILED, message[:500], None

	_send_password(row, password)
	return CREATED, row.message, user.name


def _send_password(row, password: str) -> None:
	"""The text, sent now from this job rather than queued, so the password is
	never written to the job queue. How it went lands on the row."""
	record = (ROW_DOCTYPE, row.name)
	if not sms.configured():
		sms._record(record, sms.OFF)
		return
	sms.deliver(row.phone, welcome_text(row.first_name, row.id_number, password), record=record)


# --------------------------------------------------------------------------
# counts, report, answers
# --------------------------------------------------------------------------


def _count(doc) -> None:
	tally = {}
	for row in doc.rows:
		tally[row.result] = tally.get(row.result, 0) + 1
	doc.total_rows = len(doc.rows)
	doc.existing_rows = tally.get(EXISTING, 0)
	doc.duplicate_rows = tally.get(DUPLICATE, 0)
	doc.error_rows = tally.get(ERROR, 0)
	doc.created_count = tally.get(CREATED, 0)
	doc.failed_count = tally.get(ROW_FAILED, 0)
	# New until run; after a run, the rows that were New then.
	doc.new_rows = tally.get(NEW, 0) + doc.created_count + doc.failed_count if doc.status in (
		RUNNING, COMPLETED, FAILED
	) else tally.get(NEW, 0)


REPORT_TABS = (
	("Errors", (ERROR, ROW_FAILED)),
	("Skipped", (EXISTING, DUPLICATE)),
	("Created", (CREATED,)),
)


def build_report(doc) -> bytes:
	"""The workbook the administrator downloads: every row that did not become
	an account, with why, in the columns MPS sent — and those that did. No
	password is in it."""
	from openpyxl import Workbook
	from openpyxl.styles import Font

	columns = []
	for row in doc.rows:
		for key in json.loads(row.source or "{}"):
			if key not in columns:
				columns.append(key)

	book = Workbook()
	book.remove(book.active)
	for title, results in REPORT_TABS:
		sheet = book.create_sheet(title)
		extra = ["Account", "Text"] if title == "Created" else []
		sheet.append(["Sheet", "Row", "Result", "Reason", "ID Type", *extra, *columns])
		for cell in sheet[1]:
			cell.font = Font(bold=True)
		for row in doc.rows:
			if row.result not in results:
				continue
			source = json.loads(row.source or "{}")
			values = [row.sheet, row.row_number, row.result, row.message, row.id_type]
			if extra:
				values += [row.user, row.sms_status]
			sheet.append(values + [source.get(key) for key in columns])
	buffer = io.BytesIO()
	book.save(buffer)
	return buffer.getvalue()


def _write_report(doc) -> None:
	old = doc.error_report
	stem = doc.source_file_name.rsplit(".", 1)[0]
	doc.db_set("error_report", _attach(doc, f"{doc.name} report - {stem}.xlsx", build_report(doc)))
	if old:
		for name in frappe.get_all("File", filters={"file_url": old, "attached_to_name": doc.name}, pluck="name"):
			frappe.delete_doc("File", name, ignore_permissions=True, force=True)


def _sheets(doc) -> list[dict]:
	out = {}
	for row in doc.rows:
		sheet = out.setdefault(row.sheet, {"sheet": row.sheet, "total": 0})
		sheet["total"] += 1
		key = row.result.lower()
		sheet[key] = sheet.get(key, 0) + 1
	return list(out.values())


ISSUES = (ERROR, ROW_FAILED, DUPLICATE, EXISTING)
SHOWN_ROWS = 500


def summary(doc) -> dict:
	earlier = frappe.get_all(
		DOCTYPE,
		filters={"file_hash": doc.file_hash, "name": ["!=", doc.name], "status": COMPLETED},
		pluck="name",
		order_by="creation desc",
		limit_page_length=1,
	)
	issues = [r for r in doc.rows if r.result in ISSUES]
	return {
		"name": doc.name,
		"status": doc.status,
		"file_name": doc.source_file_name,
		"uploaded_by": doc.uploaded_by,
		"uploaded_on": doc.uploaded_on,
		"run_by": doc.run_by,
		"reason": doc.reason,
		"started_on": doc.started_on,
		"finished_on": doc.finished_on,
		"counts": {
			"total": cint(doc.total_rows),
			"new": cint(doc.new_rows),
			"existing": cint(doc.existing_rows),
			"duplicate": cint(doc.duplicate_rows),
			"error": cint(doc.error_rows),
			"created": cint(doc.created_count),
			"failed": cint(doc.failed_count),
			"sms_sent": sum(1 for r in doc.rows if r.sms_status == sms.SENT),
		},
		"sheets": _sheets(doc),
		"notes": [n for n in (doc.notes or "").split("\n") if n],
		"failure": doc.failure,
		"error_report": doc.error_report,
		"already_imported_as": earlier[0] if earlier else None,
		"sms_configured": sms.configured(),
		"keycloak_configured": bool(keycloak_admin.citizen_config()),
		"issues": [
			{
				"sheet": r.sheet,
				"row_number": r.row_number,
				"result": r.result,
				"message": r.message,
				"full_name": r.full_name,
				"id_number": r.id_number,
				"id_type": r.id_type,
			}
			for r in issues[:SHOWN_ROWS]
		],
		"issues_truncated": len(issues) > SHOWN_ROWS,
	}


def get(batch: str) -> dict:
	return summary(_load(batch))


def history(start=0, page_length=20) -> dict:
	page_length = max(1, min(cint(page_length) or 20, 100))
	start = max(0, cint(start))
	rows = frappe.get_all(
		DOCTYPE,
		fields=[
			"name", "status", "source_file_name", "uploaded_by", "uploaded_on", "run_by", "finished_on",
			"total_rows", "new_rows", "existing_rows", "duplicate_rows", "error_rows", "created_count", "failed_count",
			"error_report",
		],
		order_by="creation desc",
		limit_start=start,
		limit_page_length=page_length + 1,
	)
	return {"rows": rows[:page_length], "has_more": len(rows) > page_length}
