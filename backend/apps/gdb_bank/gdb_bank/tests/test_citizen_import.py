"""Citizens from the MPS call list (gdb_bank/services/citizen_import.py).

The cleaning rules are plain functions and are tested on their own. The
workbook path is tested end to end on a workbook built here — never on a real
MPS file, which holds real people. Keycloak and SMS are stubbed: these tests
are about what GDB decides and what it writes.

    bench --site gdb.localhost run-tests --module gdb_bank.tests.test_citizen_import
"""

import io
import json
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase
from openpyxl import Workbook

from gdb_bank import tin_auth
from gdb_bank.security import role_policy
from gdb_bank.services import access_audit
from gdb_bank.services import citizen_import as ci

ADMIN = "test-gdb-import-admin@example.gy"
REASON = "Test: MPS call list"
HEADER = ["NAME", "ADDRESS", "REGION", "CONTACT #", "BUSINESS", "TYPE OF BUSINESS", "ID Number", "Amount Req", "Outcome of Call"]


def _workbook(*sheets: tuple[str, list[list]]) -> bytes:
	"""A workbook laid out as MPS sends it: a title, a date, then the header."""
	book = Workbook()
	book.remove(book.active)
	for title, rows in sheets:
		sheet = book.create_sheet(title)
		sheet.append(["Successful Calls for Directly routing to GDB"])
		sheet.append(["6th Oct 2026"])
		sheet.append(HEADER)
		for row in rows:
			sheet.append(row)
	buffer = io.BytesIO()
	book.save(buffer)
	return buffer.getvalue()


def _row(name="Asha Test Persaud", id_number=900000001.0, phone="600-1001", region=3.0, address="1 Main Street"):
	return [name, address, region, phone, "New", "Poultry", id_number, 3000000.0, "Route to GDB"]


class TestCleaning(IntegrationTestCase):
	def test_the_last_word_is_the_last_name(self):
		self.assertEqual(ci.split_name("Asha Persaud"), ("Asha", "Persaud"))
		self.assertEqual(ci.split_name("  Asha   Devi Persaud "), ("Asha Devi", "Persaud"))
		self.assertEqual(ci.split_name("Mary-Ann O'Neil"), ("Mary-Ann", "O'Neil"))

	def test_a_name_must_be_a_name(self):
		for bad in ("", None, "Cher", "Ravi Singh - +1 (555) 010-0000"):
			with self.assertRaises(ci.RowError):
				ci.split_name(bad)

	def test_an_id_beginning_with_a_letter_is_a_passport(self):
		self.assertEqual(ci.classify_id("R9000001"), (ci.PASSPORT, "R9000001"))
		self.assertEqual(ci.classify_id("re 9000002"), (ci.PASSPORT, "RE9000002"))

	def test_nine_digits_is_a_national_id_and_eleven_an_eid(self):
		self.assertEqual(ci.classify_id(900000003.0), (ci.NATIONAL_ID, "900000003"))
		self.assertEqual(ci.classify_id("900000003.0"), (ci.NATIONAL_ID, "900000003"))
		self.assertEqual(ci.classify_id("592-2001-0101"), (ci.EID, "59220010101"))

	def test_anything_else_is_an_error(self):
		for bad in (None, "", "NA", 12345678.0, 1234567890.0):
			with self.assertRaises(ci.RowError):
				ci.classify_id(bad)

	def test_phones_take_the_first_number_as_592(self):
		self.assertEqual(ci.clean_phone("600-2995"), "+5926002995")
		self.assertEqual(ci.clean_phone(6002995.0), "+5926002995")
		self.assertEqual(ci.clean_phone("600 2995"), "+5926002995")
		self.assertEqual(ci.clean_phone("601-8294 / 655-1234"), "+5926018294")
		with self.assertRaises(ci.RowError):
			ci.clean_phone("")

	def test_regions_map_onto_the_list_and_never_stop_a_row(self):
		self.assertEqual(ci.clean_region(2.0), ("Region 2 — Pomeroon-Supenaam", None))
		self.assertEqual(ci.clean_region("Region 10"), ("Region 10 — Upper Demerara-Berbice", None))
		region, warning = ci.clean_region(14)
		self.assertIsNone(region)
		self.assertIn("left blank", warning)

	def test_every_problem_in_a_row_is_reported(self):
		row = ci.clean_row({"name": "Cher", "id_number": "NA", "phone": "", "region": 2, "address": ""})
		self.assertEqual(row["result"], ci.ERROR)
		for part in ("one word", "NA", "Contact number is blank", "Address is blank"):
			self.assertIn(part, row["message"])

	def test_the_text_is_at_most_150_characters(self):
		password = ci.temporary_password()
		short = ci.welcome_text("Asha", "900000003", password)
		self.assertTrue(short.startswith("Hi Asha, your GDB account is created via MPSGEI."))
		self.assertIn("900000003", short)
		self.assertIn(password, short)
		self.assertLessEqual(len(short), ci.SMS_MAX)
		long = ci.welcome_text("Bartholomewandersonius Jane", "RE90000021234", password)
		self.assertLessEqual(len(long), ci.SMS_MAX)
		self.assertIn("RE90000021234", long)
		self.assertIn(password, long)

	def test_temporary_passwords_mix_cases_and_digits(self):
		for _ in range(20):
			password = ci.temporary_password()
			self.assertEqual(len(password), ci.PASSWORD_LENGTH)
			self.assertTrue(any(c.isupper() for c in password) and any(c.islower() for c in password))
			self.assertTrue(any(c.isdigit() for c in password))


class ImportCase(IntegrationTestCase):
	def setUp(self):
		super().setUp()
		self.enterContext(patch.object(frappe.local.db, "commit"))
		self.addCleanup(frappe.db.rollback)
		# This site may already have made an hour's worth of Users (a real
		# import does); the test's own fixtures must not be throttled by that.
		self.enterContext(ci._room_for_new_users())
		frappe.get_doc(
			{
				"doctype": "User",
				"email": ADMIN,
				"first_name": "Import Admin",
				"user_type": "System User",
				"send_welcome_email": 0,
				"roles": [{"role": "Platform Admin"}],
			}
		).insert(ignore_permissions=True)
		self.create = self.enterContext(
			patch.object(ci.keycloak_admin, "create_citizen_account", return_value="kc-id")
		)
		self.delete = self.enterContext(patch.object(ci.keycloak_admin, "delete_citizen_account"))
		self.enterContext(patch.object(ci.keycloak_admin, "citizen_config", return_value={"realm": "gdb-citizen"}))
		self.enterContext(patch.object(ci.sms, "configured", return_value=True))
		self.deliver = self.enterContext(patch.object(ci.sms, "deliver", return_value={"sent": True}))
		self.enqueue = self.enterContext(patch.object(ci.frappe, "enqueue"))
		frappe.set_user(ADMIN)
		self.addCleanup(frappe.set_user, "Administrator")

	def preview(self, content: bytes) -> dict:
		return ci.preview(ADMIN, "GDP List.xlsx", content)

	def run_import(self, batch: str) -> dict:
		ci.run(ADMIN, batch, REASON)
		ci.execute(batch)
		return ci.get(batch)


class TestPreview(ImportCase):
	def test_rows_are_sorted_and_nothing_is_created(self):
		content = _workbook(
			("R1", [_row(), _row(name="Cher", id_number="NA"), _row(name="Bibi Khan", id_number="R9000001", phone="601-2222")]),
			("R2", [_row(name="Asha Again", phone="602-3333")]),
		)
		result = self.preview(content)
		self.assertEqual(result["status"], ci.PREVIEWED)
		self.assertEqual(result["counts"]["total"], 4)
		self.assertEqual(result["counts"]["new"], 2)
		self.assertEqual(result["counts"]["error"], 1)
		self.assertEqual(result["counts"]["duplicate"], 1)
		self.assertTrue(result["error_report"])
		self.create.assert_not_called()
		self.assertFalse(frappe.db.exists("User", {tin_auth.NID_FIELD: "900000001"}))

	def test_columns_are_found_by_name_not_position(self):
		book = Workbook()
		sheet = book.active
		sheet.title = "Reordered"
		sheet.append(["ID Number", "CONTACT #", "NAME", "REGION", "ADDRESS"])
		sheet.append([900000002, "600-1002", "Ravi Test Singh", 4, "Lot 2"])
		buffer = io.BytesIO()
		book.save(buffer)
		result = self.preview(buffer.getvalue())
		self.assertEqual(result["counts"]["new"], 1)

	def test_a_sheet_without_the_columns_is_skipped_and_said(self):
		book = Workbook()
		book.active.append(["NAME", "ADDRESS"])
		book.active.append(["Ravi Test Singh", "Lot 2"])
		good = book.create_sheet("Good")
		good.append(HEADER)
		good.append(_row())
		buffer = io.BytesIO()
		book.save(buffer)
		result = self.preview(buffer.getvalue())
		self.assertEqual(result["counts"]["total"], 1)
		self.assertTrue(any("region" in note for note in result["notes"]))

	def test_an_id_an_account_already_holds_is_skipped(self):
		frappe.get_doc(
			{
				"doctype": "User",
				"email": "test-gdb-existing@tin.gdb.invalid",
				"first_name": "Existing",
				"user_type": "Website User",
				"send_welcome_email": 0,
				tin_auth.NID_FIELD: "900000001",
			}
		).insert(ignore_permissions=True)
		result = self.preview(_workbook(("R1", [_row()])))
		self.assertEqual(result["counts"]["existing"], 1)
		self.assertEqual(result["counts"]["new"], 0)

	def test_only_a_real_xlsx_is_read(self):
		with self.assertRaisesRegex(frappe.ValidationError, "not an Excel"):
			self.preview(b"NAME,ADDRESS\nAsha,Lot 1\n")


class TestRun(ImportCase):
	def test_each_new_row_becomes_a_citizen_and_is_texted(self):
		batch = self.preview(_workbook(("R1", [_row(), _row(name="Bibi Khan", id_number="R9000001", phone="601-2222")])))["name"]
		result = self.run_import(batch)

		self.assertEqual(result["status"], ci.COMPLETED)
		self.assertEqual(result["counts"]["created"], 2)
		user = frappe.db.get_value("User", {tin_auth.NID_FIELD: "900000001"}, ["name", "first_name", "last_name", "user_type", "mobile_no", ci.ID_TYPE_FIELD], as_dict=True)
		self.assertEqual(user.name, "900000001@tin.gdb.invalid")
		self.assertEqual((user.first_name, user.last_name), ("Asha Test", "Persaud"))
		self.assertEqual(user.user_type, "Website User")
		self.assertEqual(user.mobile_no, "+5926001001")
		self.assertEqual(user[ci.ID_TYPE_FIELD], ci.NATIONAL_ID)
		self.assertEqual(frappe.get_roles(user.name).count("Citizen"), 1)
		self.assertEqual(frappe.db.get_value("User", {tin_auth.NID_FIELD: "R9000001"}, ci.ID_TYPE_FIELD), ci.PASSPORT)

		profile = frappe.db.get_value(
			"GDB Citizen Profile", {"user": user.name}, ["phone", "national_id", "region", "address", "declared_source"], as_dict=True
		)
		self.assertEqual(profile.national_id, "900000001")
		self.assertEqual(profile.region, "Region 3 — Essequibo Islands-West Demerara")
		self.assertEqual(profile.address, "1 Main Street")
		self.assertIn(batch, profile.declared_source)

		# Keycloak: the ID as username, a placeholder email, a TEMPORARY password.
		args, kwargs = self.create.call_args_list[0]
		self.assertEqual(args[:4], ("900000001", "900000001@tin.gdb.invalid", "Asha Test", "Persaud"))
		self.assertTrue(kwargs["temporary"])
		password = args[4]
		# The same password, texted to the cleaned phone, and kept nowhere else.
		to, body = self.deliver.call_args_list[0].args
		self.assertEqual(to, "+5926001001")
		self.assertIn(password, body)
		self.assertLessEqual(len(body), ci.SMS_MAX)
		doc = frappe.get_doc(ci.DOCTYPE, batch)
		self.assertNotIn(password, json.dumps(doc.as_dict(), default=str))

		actions = [r.action for r in frappe.get_all(access_audit.DOCTYPE, filters={"subject_user": user.name}, fields=["action"])]
		self.assertIn(access_audit.ACCOUNT_CREATED, actions)
		self.assertIn(access_audit.ONE_TIME_PASSWORD_ISSUED, actions)

	def test_a_batch_runs_once(self):
		batch = self.preview(_workbook(("R1", [_row()])))["name"]
		self.run_import(batch)
		again = ci.run(ADMIN, batch, REASON)
		self.assertEqual(again["status"], ci.COMPLETED)
		queued = [c for c in self.enqueue.call_args_list if c.args and c.args[0].endswith("citizen_import.execute")]
		self.assertEqual(len(queued), 1)
		self.assertEqual(self.create.call_count, 1)

	def test_a_reason_is_required(self):
		batch = self.preview(_workbook(("R1", [_row()])))["name"]
		with self.assertRaisesRegex(frappe.ValidationError, "Say why"):
			ci.run(ADMIN, batch, " ")

	def test_keycloak_already_holding_the_id_is_skipped_not_failed(self):
		self.create.side_effect = ci.keycloak_admin.AccountExists("exists")
		batch = self.preview(_workbook(("R1", [_row()])))["name"]
		result = self.run_import(batch)
		self.assertEqual(result["counts"]["existing"], 1)
		self.assertEqual(result["counts"]["created"], 0)
		self.deliver.assert_not_called()

	def test_a_keycloak_failure_fails_the_row_and_the_rest_go_on(self):
		self.create.side_effect = [ci.keycloak_admin.KeycloakAdminError("down"), "kc-2"]
		batch = self.preview(_workbook(("R1", [_row(), _row(name="Bibi Khan", id_number=900000002, phone="601-2222")])))["name"]
		result = self.run_import(batch)
		self.assertEqual(result["counts"]["failed"], 1)
		self.assertEqual(result["counts"]["created"], 1)
		self.assertEqual(result["status"], ci.COMPLETED)

	def test_frappes_hourly_user_throttle_does_not_stop_the_import(self):
		# More than an hour's worth of new Users already: Frappe would say
		# "Throttled" to any other caller.
		before = frappe.local.conf.get("throttle_user_limit")
		limit = int(before or 60)
		batch = self.preview(_workbook(("R1", [_row()])))["name"]
		with patch.object(ci.frappe.db, "get_creation_count", return_value=limit + 50):
			result = self.run_import(batch)
		self.assertEqual(result["counts"]["created"], 1)
		self.assertEqual(result["counts"]["failed"], 0)
		# ...and the guard is back as it was for everything after the job.
		self.assertEqual(frappe.local.conf.get("throttle_user_limit"), before)

	def test_without_sms_the_account_is_still_made_and_the_row_says_so(self):
		with patch.object(ci.sms, "configured", return_value=False):
			batch = self.preview(_workbook(("R1", [_row()])))["name"]
			self.run_import(batch)
		self.deliver.assert_not_called()
		row = frappe.get_all(ci.ROW_DOCTYPE, filters={"parent": batch}, fields=["result", "sms_status"])[0]
		self.assertEqual(row.result, ci.CREATED)
		self.assertEqual(row.sms_status, ci.sms.OFF)


class TestThePolicy(ImportCase):
	def test_an_admin_may_open_a_citizen_account_and_nothing_more(self):
		self.assertTrue(role_policy.binds(ADMIN))
		frappe.get_doc(
			{
				"doctype": "User",
				"email": "test-gdb-citizen-ok@example.gy",
				"first_name": "Allowed",
				"user_type": "Website User",
				"send_welcome_email": 0,
				"roles": [{"role": "Citizen"}],
			}
		).insert(ignore_permissions=True)
		with self.assertRaises(frappe.PermissionError):
			frappe.get_doc(
				{
					"doctype": "User",
					"email": "test-gdb-sneaky@example.gy",
					"first_name": "Sneaky",
					"user_type": "Website User",
					"send_welcome_email": 0,
					"roles": [{"role": "Citizen"}, {"role": "Loan Underwriter"}],
				}
			).insert(ignore_permissions=True)


class TestFirstSignIn(IntegrationTestCase):
	"""tin_auth.set_initial_password, with Keycloak stubbed."""

	def setUp(self):
		super().setUp()
		import inspect

		self.enterContext(patch.object(frappe.local.db, "commit"))
		self.addCleanup(frappe.db.rollback)
		# This site may already have made an hour's worth of Users (a real
		# import does); the test's own fixtures must not be throttled by that.
		self.enterContext(ci._room_for_new_users())
		self.enterContext(patch.object(tin_auth, "_require_door"))
		self.enterContext(patch.object(tin_auth.sms_otp, "configured", return_value=False))
		self.enterContext(patch.object(tin_auth.identity, "keycloak_settings", return_value={"population": "citizen", "realm": "gdb-citizen"}))
		self.enterContext(patch.object(tin_auth.keycloak_admin, "citizen_user_id", return_value="kc-id"))
		self.set_password = self.enterContext(patch.object(tin_auth.keycloak_admin, "set_citizen_password"))
		self.login = inspect.unwrap(tin_auth.national_id_login)
		self.first = inspect.unwrap(tin_auth.set_initial_password)
		frappe.get_doc(
			{
				"doctype": "User",
				"email": "900000009@tin.gdb.invalid",
				"first_name": "First",
				"last_name": "Signin",
				"mobile_no": "+5926009999",
				"user_type": "Website User",
				"send_welcome_email": 0,
				tin_auth.NID_FIELD: "900000009",
				"roles": [{"role": "Citizen"}],
			}
		).insert(ignore_permissions=True)

	def test_a_temporary_password_asks_for_a_new_one(self):
		with patch.object(tin_auth.identity, "_request_token", side_effect=tin_auth.identity._PasswordChangeRequired):
			self.assertEqual(self.login("900000009", "Tmp2abcd"), {"password_change_required": True})

	def test_the_new_password_is_set_then_a_code_is_sent(self):
		with patch.object(tin_auth.identity, "_request_token", side_effect=tin_auth.identity._PasswordChangeRequired):
			out = self.first("900000009", "Tmp2abcd", "Garden2026", "Garden2026")
		self.set_password.assert_called_once_with("kc-id", "Garden2026")
		self.assertTrue(out["otp_required"])
		self.assertTrue(out["challenge"])

	def test_an_account_not_waiting_on_a_password_is_refused(self):
		with patch.object(tin_auth.identity, "_request_token", return_value="token"):
			with self.assertRaisesRegex(frappe.ValidationError, "already set"):
				self.first("900000009", "Garden2026", "Orchard2027", "Orchard2027")
		with patch.object(tin_auth.identity, "_request_token", return_value=None):
			with self.assertRaises(frappe.AuthenticationError):
				self.first("900000009", "wrong", "Orchard2027", "Orchard2027")
		self.set_password.assert_not_called()

	def test_the_new_password_follows_the_rules(self):
		with patch.object(tin_auth.identity, "_request_token", side_effect=tin_auth.identity._PasswordChangeRequired):
			with self.assertRaisesRegex(frappe.ValidationError, "own"):
				self.first("900000009", "Tmp2abcd9", "Tmp2abcd9", "Tmp2abcd9")
			with self.assertRaisesRegex(frappe.ValidationError, "letter and one number"):
				self.first("900000009", "Tmp2abcd", "gardengarden", "gardengarden")
